"""Synthetic file-behavior checks; no course content or persistent PPTX fixtures."""

import contextlib
import importlib.util
import io
import os
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location(
    "audit_pptx", ROOT / "skills/ielts-speaking-workflow/scripts/audit_pptx.py")
AUDITOR = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(AUDITOR)


class AuditTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "synthetic.pptx"
        self.presentation = Presentation()

    def add_slide(self):
        return self.presentation.slides.add_slide(self.presentation.slide_layouts[6])

    def save_and_audit(self):
        self.presentation.save(self.path)
        before = self.path.read_bytes()
        report = AUDITOR.audit(self.path)
        self.assertEqual(before, self.path.read_bytes())
        return report

    def test_no_visual_or_font_family_preferences(self):
        slide = self.add_slide()
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "Text"
        run.font.name = "Aptos"
        run.font.size = Pt(28)
        report = self.save_and_audit()
        self.assertEqual([], report["findings"])

    def test_off_canvas_and_small_text(self):
        slide = self.add_slide()
        box = slide.shapes.add_textbox(Inches(20), Inches(1), Inches(2), Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "Text"
        run.font.size = Pt(8)
        report = self.save_and_audit()
        self.assertGreater(report["counts"]["error"], 0)
        self.assertGreater(report["counts"]["advisory"], 0)

    def test_notes_and_hidden_flags_follow_display_order(self):
        first, second = self.add_slide(), self.add_slide()
        first.notes_slide.notes_text_frame.text = "Private note"
        second._element.set("show", "0")
        order = self.presentation.slides._sldIdLst
        order.insert(0, order[1])
        report = self.save_and_audit()
        self.assertEqual([2], report["metadata"]["notes"])
        self.assertEqual([1], report["metadata"]["hidden_slides"])
        self.assertEqual(0, report["counts"]["error"])
        self.assertNotIn("Private note", str(report))

    def test_automatically_inserted_notes_page_number_is_not_note_content(self):
        slide = self.add_slide()
        notes = slide.notes_slide
        # Default notes pages have slide-number placeholders, not teacher prose.
        self.assertFalse(notes.notes_text_frame.text.strip())
        report = self.save_and_audit()
        self.assertEqual([], report["metadata"]["notes"])

    def test_nested_and_table_text_is_read(self):
        slide = self.add_slide()
        group = slide.shapes.add_group_shape()
        textbox = group.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
        textbox.text = "Text"
        textbox.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
        table = slide.shapes.add_table(1, 1, Inches(3), Inches(1), Inches(2), Inches(1)).table
        table.cell(0, 0).text = "Text"
        table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.size = Pt(8)
        report = self.save_and_audit()
        small = [item for item in report["findings"] if item["id"] == "ppt/small-explicit-text"]
        self.assertEqual(2, len(small))

    def test_reports_and_error_exit_without_modifying_input(self):
        slide = self.add_slide()
        box = slide.shapes.add_textbox(Inches(20), Inches(1), Inches(2), Inches(1))
        box.text = "Text"
        self.presentation.save(self.path)
        before = self.path.read_bytes()
        json_file, md_file = self.path.with_suffix(".json"), self.path.with_suffix(".md")
        with contextlib.redirect_stdout(io.StringIO()):
            exit_code = AUDITOR.main([str(self.path), "--json", str(json_file),
                                      "--markdown", str(md_file), "--fail-on-error"])
        self.assertEqual(1, exit_code)
        self.assertEqual(before, self.path.read_bytes())
        self.assertTrue(json_file.is_file())
        self.assertTrue(md_file.is_file())

    def test_report_cannot_overwrite_input_or_its_hardlink(self):
        self.add_slide()
        self.presentation.save(self.path)
        before = self.path.read_bytes()
        linked = self.path.with_name("report.json")
        os.link(self.path, linked)
        for output in (self.path, linked):
            with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit) as result:
                AUDITOR.main([str(self.path), "--json", str(output)])
            self.assertEqual(2, result.exception.code)
        self.assertEqual(before, self.path.read_bytes())

    def test_invalid_file_returns_failure(self):
        self.path.write_bytes(b"not a PPTX")
        with contextlib.redirect_stderr(io.StringIO()):
            exit_code = AUDITOR.main([str(self.path)])
        self.assertEqual(2, exit_code)
        self.assertEqual(b"not a PPTX", self.path.read_bytes())


if __name__ == "__main__":
    unittest.main()
