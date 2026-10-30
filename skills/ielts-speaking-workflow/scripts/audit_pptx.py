#!/usr/bin/env python3
"""Read-only PPTX triage / 只读 PPTX 初检；does not render or edit slides."""

from __future__ import annotations

import argparse
import json
import math
import posixpath
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError as exc:
    raise SystemExit("Install the optional dependency: python -m pip install python-pptx==1.0.2") from exc


NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def finding(rule: str, severity: str, slide: int, evidence: str,
            recommendation: str, shape: str | None = None) -> dict[str, Any]:
    result = {"id": rule, "severity": severity, "slide": slide,
              "evidence": evidence, "recommendation": recommendation}
    if shape:
        result["shape"] = shape
    return result


def part_target(source: str, target: str) -> str:
    return posixpath.normpath(posixpath.join(posixpath.dirname(source), target)).lstrip("/")


def package_metadata(path: Path) -> dict[str, Any]:
    """Resolve slide relationships in display order, not slide XML filename order."""
    result: dict[str, Any] = {"notes": [], "hidden_slides": [], "comments": {},
                              "hidden_shapes": {}, "unlinked_comment_parts": 0}
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())

        def relationships(source: str) -> list[ET.Element]:
            folder, basename = posixpath.split(source)
            relpath = posixpath.join(folder, "_rels", basename + ".rels")
            return list(ET.fromstring(archive.read(relpath))) if relpath in names else []

        presentation = ET.fromstring(archive.read("ppt/presentation.xml"))
        rels = {rel.get("Id"): rel for rel in relationships("ppt/presentation.xml")}
        ids = presentation.find("p:sldIdLst", NS)
        comment_parts = {
            name for name in names if name.endswith(".xml") and
            (name.startswith("ppt/comments/") or name.startswith("ppt/threadedComments/"))
        }
        linked_comments: set[str] = set()
        for number, slide_id in enumerate(list(ids) if ids is not None else [], 1):
            rel = rels.get(slide_id.get("{" + NS["r"] + "}id"))
            if rel is None or rel.get("TargetMode") == "External":
                raise ValueError(f"Invalid internal slide relationship at displayed slide {number}")
            slide_part = part_target("ppt/presentation.xml", rel.get("Target", ""))
            root = ET.fromstring(archive.read(slide_part))
            if any(node.get("show", "1").lower() in {"0", "false"} for node in (slide_id, root)):
                result["hidden_slides"].append(number)
            hidden = sum(1 for node in root.iter() if node.get("hidden", "0").lower() in {"1", "true"})
            if hidden:
                result["hidden_shapes"][number] = hidden
            for child_rel in relationships(slide_part):
                if child_rel.get("TargetMode") == "External":
                    continue
                target = part_target(slide_part, child_rel.get("Target", ""))
                relation_type = child_rel.get("Type", "").rsplit("/", 1)[-1]
                if relation_type == "notesSlide":
                    notes = ET.fromstring(archive.read(target))
                    texts = []
                    for shape in notes.findall(".//p:sp", NS):
                        placeholder = shape.find(".//p:ph", NS)
                        if placeholder is not None and placeholder.get("type") in {"sldNum", "hdr", "ftr", "dt"}:
                            continue
                        texts.extend(node.text or "" for node in shape.findall(".//a:t", NS))
                    if "".join(texts).strip():
                        result["notes"].append(number)
                elif relation_type in {"comments", "threadedComment", "threadedComments"}:
                    linked_comments.add(target)
                    comments = ET.fromstring(archive.read(target))
                    count = sum(1 for node in comments.iter()
                                if node.tag.rsplit("}", 1)[-1] in {"cm", "comment", "threadedComment"})
                    result["comments"][number] = result["comments"].get(number, 0) + max(1, count)
        result["unlinked_comment_parts"] = len(comment_parts - linked_comments)
    return result


def text_frames(shapes: Any):
    """Read group/table text, without assuming their local geometry is slide geometry."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from text_frames(shape.shapes)
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    yield shape.name, cell.text_frame
        elif getattr(shape, "has_text_frame", False):
            yield shape.name, shape.text_frame


def estimated_text_height(shape: Any) -> tuple[float, float]:
    """Rough text-fit signal only; theme fonts and wrapping require a renderer."""
    frame = shape.text_frame
    width = max(1, (shape.width - frame.margin_left - frame.margin_right) / 12700)
    height = max(0, (shape.height - frame.margin_top - frame.margin_bottom) / 12700)
    sizes = [run.font.size.pt for paragraph in frame.paragraphs for run in paragraph.runs
             if run.text and run.font.size is not None]
    size = max(sizes, default=20.0)
    lines = 0
    for line in frame.text.split("\n"):
        units = sum(1.0 if "\u3400" <= char <= "\u9fff" else 0.55 for char in line)
        lines += max(1, math.ceil(units * size / width))
    return lines * size * 1.2, height


def audit(path: Path, min_font_size: float = 14.0) -> dict[str, Any]:
    presentation = Presentation(str(path))
    metadata = package_metadata(path)
    findings: list[dict[str, Any]] = []
    summaries = []
    width, height = presentation.slide_width, presentation.slide_height
    for number, slide in enumerate(presentation.slides, 1):
        shapes = list(slide.shapes)
        for name, frame in text_frames(shapes):
            small = sorted({round(run.font.size.pt, 1) for paragraph in frame.paragraphs
                            for run in paragraph.runs if run.text.strip() and
                            run.font.size is not None and run.font.size.pt < min_font_size})
            if small:
                findings.append(finding("ppt/small-explicit-text", "advisory", number,
                                        f"Explicit type sizes below {min_font_size:g} pt: {small}",
                                        "Check projection readability; theme-inherited sizes are not evaluated.", name))
        for shape in shapes:
            has_text = getattr(shape, "has_text_frame", False) and bool(shape.text.strip())
            is_picture = shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            if not (has_text or is_picture):
                continue
            left, top = shape.left, shape.top
            right, bottom = left + shape.width, top + shape.height
            outside = right <= 0 or bottom <= 0 or left >= width or top >= height
            partial = left < 0 or top < 0 or right > width or bottom > height
            if outside or partial:
                findings.append(finding("ppt/content-outside-canvas", "error" if outside else "warning", number,
                                        "Content is wholly outside the slide" if outside else "Content crosses a slide edge",
                                        "Inspect intended cropping or bleed; move unintended off-slide content.", shape.name))
            if has_text and len(shape.text.strip()) > 25:
                estimated, available = estimated_text_height(shape)
                if estimated > available * 1.12:
                    findings.append(finding("ppt/text-fit-estimate", "warning", number,
                                            f"Estimated text height {estimated:.1f} pt; available {available:.1f} pt",
                                            "Render to confirm; enlarge the box or simplify text if clipping is visible.", shape.name))
        for index, first in enumerate(shapes):
            first_text = getattr(first, "has_text_frame", False) and bool(first.text.strip())
            if not first_text:
                continue
            for second in shapes[index + 1:]:
                second_text = getattr(second, "has_text_frame", False) and bool(second.text.strip())
                if not (second_text or second.shape_type == MSO_SHAPE_TYPE.PICTURE):
                    continue
                overlap_w = min(first.left + first.width, second.left + second.width) - max(first.left, second.left)
                overlap_h = min(first.top + first.height, second.top + second.height) - max(first.top, second.top)
                smaller = min(first.width * first.height, second.width * second.height)
                if overlap_w > 0 and overlap_h > 0 and smaller and overlap_w * overlap_h / smaller > 0.35:
                    findings.append(finding("ppt/possible-occlusion", "warning", number,
                                            "Text intersects another text object or an image above it",
                                            "Inspect z-order and transparency in a render; overlap can be intentional.",
                                            f"{first.name} / {second.name}"))
        for key, rule, message in [
            ("notes", "ppt/speaker-notes", "Non-empty speaker notes"),
            ("hidden_slides", "ppt/hidden-slide", "Slide marked hidden"),
            ("comments", "ppt/comments", "Embedded comments"),
            ("hidden_shapes", "ppt/hidden-shapes", "Objects marked hidden"),
        ]:
            if number in metadata[key]:
                findings.append(finding(rule, "advisory", number, message,
                                        "Confirm that embedded content is appropriate for the intended recipients."))
        summaries.append({"slide": number, "top_level_shapes": len(shapes),
                          "has_notes": number in metadata["notes"],
                          "hidden": number in metadata["hidden_slides"]})
    if metadata["unlinked_comment_parts"]:
        findings.append(finding("ppt/unlinked-comments", "advisory", 0,
                                f"Unlinked comment parts: {metadata['unlinked_comment_parts']}",
                                "Review embedded package content before sharing."))
    return {"file": path.name, "slide_count": len(presentation.slides),
            "counts": {severity: sum(item["severity"] == severity for item in findings)
                       for severity in ("error", "warning", "advisory")},
            "findings": findings, "metadata": metadata, "slides": summaries,
            "limits": ["No rendering or pedagogical assessment", "Text fit and overlap are heuristic",
                       "Theme fonts, rotated text, group/table geometry and animations require visual review"]}


def markdown(report: dict[str, Any]) -> str:
    lines = [f"# PPTX audit / PPTX 初检: {report['file']}", "",
             f"Slides / 页数: {report['slide_count']}", ""]
    if not report["findings"]:
        lines.append("No mechanical findings. Rendering remains unverified. / 未发现机械问题，仍需渲染复核。")
    for item in report["findings"]:
        location = f"slide {item['slide']}" if item["slide"] else "deck"
        lines.extend([f"- **{item['severity']}** `{item['id']}` ({location})",
                      f"  - {item['evidence']}", f"  - {item['recommendation']}"])
    lines.extend(["", "## Limits / 限制", ""])
    lines.extend(f"- {limit}" for limit in report["limits"])
    return "\n".join(lines) + "\n"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--json", type=Path, help="write a JSON report / 输出 JSON 报告")
    parser.add_argument("--markdown", type=Path, help="write a Markdown report / 输出 Markdown 报告")
    parser.add_argument("--min-font-size", type=float, default=14.0, help="explicit-size advisory threshold; default 14 pt")
    parser.add_argument("--fail-on-error", action="store_true", help="return 1 for error findings")
    args = parser.parse_args(argv)
    outputs = [path for path in (args.json, args.markdown) if path is not None]
    if not args.pptx.is_file():
        parser.error("input file does not exist")
    if not math.isfinite(args.min_font_size) or args.min_font_size <= 0:
        parser.error("--min-font-size must be a positive finite number")
    resolved = [path.resolve() for path in outputs]
    if args.pptx.resolve() in resolved or any(path.exists() and path.samefile(args.pptx) for path in outputs):
        parser.error("a report must not overwrite the source PPTX")
    if len(set(resolved)) != len(resolved):
        parser.error("report paths must be different")
    if any(not path.parent.is_dir() for path in outputs):
        parser.error("report parent directories must already exist")
    try:
        report = audit(args.pptx, args.min_font_size)
        json_text = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
        if args.json:
            args.json.write_text(json_text, encoding="utf-8")
        if args.markdown:
            args.markdown.write_text(markdown(report), encoding="utf-8")
        print(json_text, end="")
    except Exception as exc:
        # python-pptx and malformed OOXML expose several different exception types.
        print(f"Cannot audit PPTX / 无法完成检查: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2
    return 1 if args.fail_on_error and report["counts"]["error"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
